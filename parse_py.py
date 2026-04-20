import openpyxl
import json

wb = openpyxl.load_workbook('Preflop V2.xlsx', data_only=True)
res = {}

for name in wb.sheetnames:
    ws = wb[name]
    blocks = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, min(ws.max_column + 1, 30)):
            cell = ws.cell(row=r, column=c)
            if cell.value and str(cell.value).strip() == 'AA':
                grid = []
                for i in range(13):
                    row_data = []
                    for j in range(13):
                        c_local = ws.cell(row=r+i, column=c+j)
                        val = str(c_local.value).strip() if c_local.value else None
                        color = None
                        if c_local.fill and hasattr(c_local.fill, 'fgColor') and c_local.fill.fgColor:
                            color = c_local.fill.fgColor.rgb
                            if type(color) != str:
                                color = str(color)
                        row_data.append({"val": val, "bg": color})
                    grid.append(row_data)
                blocks.append({"r": r, "c": c, "grid": grid})
    res[name] = blocks

with open('parsed_openpyxl.json', 'w', encoding='utf-8') as f:
    json.dump(res, f, indent=2)

print("Done openpyxl")
